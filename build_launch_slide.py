from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BG        = RGBColor(0x0D, 0x2C, 0x2C)
TEAL      = RGBColor(0x4C, 0xC8, 0xC3)
TEAL_DIM  = RGBColor(0x28, 0x80, 0x7D)
TEAL_CARD = RGBColor(0x12, 0x3A, 0x3A)
TEAL_DARK = RGBColor(0x0E, 0x28, 0x28)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE  = RGBColor(0xE8, 0xF4, 0xF3)
GREY      = RGBColor(0x88, 0xAA, 0xA8)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

blank_layout = prs.slide_layouts[6]
sl = prs.slides.add_slide(blank_layout)

# ── Background ──────────────────────────────────────────────────────────────
bg = sl.background
bg.fill.solid()
bg.fill.fore_color.rgb = BG

def box(x, y, w, h, color):
    s = sl.shapes.add_shape(1, x, y, w, h)
    s.line.fill.background()
    s.fill.solid()
    s.fill.fore_color.rgb = color
    return s

def rbox(x, y, w, h, color):
    s = sl.shapes.add_shape(5, x, y, w, h)  # rounded rect
    s.line.fill.background()
    s.fill.solid()
    s.fill.fore_color.rgb = color
    try: s.adjustments[0] = 0.06
    except: pass
    return s

def oval(x, y, w, h, color, line_color=None, line_w=None):
    s = sl.shapes.add_shape(9, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = Pt(line_w or 2)
    else:
        s.line.fill.background()
    return s

def tb(text, x, y, w, h, size=14, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, italic=False):
    t = sl.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return t

def tb2(lines, x, y, w, h, align=PP_ALIGN.LEFT):
    """lines = list of (text, size, bold, color, italic, space_before)"""
    t = sl.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame
    tf.word_wrap = True
    first = True
    for text, size, bold, color, italic, space_before in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if space_before:
            p.space_before = Pt(space_before)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return t

# ── Top accent bar ───────────────────────────────────────────────────────────
box(0, Inches(0.0), W, Inches(0.055), TEAL)

# ── Header labels ────────────────────────────────────────────────────────────
tb("TEST PLAN + ASK", Inches(0.55), Inches(0.18), Inches(4.0), Inches(0.38),
   size=11, bold=True, color=TEAL)

tb("90-Day Launch Plan", Inches(0.55), Inches(0.58), Inches(8.5), Inches(1.0),
   size=52, bold=True, color=WHITE)

tb("Build it. Build the audience. Drop it.", Inches(0.55), Inches(1.55),
   Inches(7.5), Inches(0.48), size=19, color=OFFWHITE)

# ── Timeline line (vertical connector) ──────────────────────────────────────
box(Inches(0.88), Inches(2.22), Inches(0.04), Inches(4.55),
    RGBColor(0x28, 0x80, 0x7D))

# ── Phase cards ──────────────────────────────────────────────────────────────
phases = [
    ("Phase 1", "Days 1–30",  "Build the Foundation",
     "Create the website, app, and core platform."),
    ("Phase 2", "Days 31–60", "Build the Audience",
     "Launch parent groups on Facebook, TikTok, and Instagram. "
     "Hire a small social team to create content, make videos, and run ads."),
    ("Phase 3", "Days 61–90", "Drop & Launch",
     "Open subscriptions, promote the launch, and send the first mail drop."),
]

card_heights = [Inches(1.05), Inches(1.38), Inches(1.05)]
card_y_start = Inches(2.1)
gap = Inches(0.18)

cy = card_y_start
for i, (phase, days, head, body) in enumerate(phases):
    ch = card_heights[i]
    # Timeline dot
    r = Inches(0.18)
    oval(Inches(0.9) - r, cy + ch/2 - r, r*2, r*2,
         TEAL_CARD, line_color=TEAL, line_w=2.5)

    # Full card background
    rbox(Inches(1.18), cy, Inches(7.55), ch, TEAL_CARD)

    # Left label panel
    rbox(Inches(1.18), cy, Inches(1.85), ch, TEAL_DIM)

    # Phase label + days
    tb2([
        (phase, 14, True, WHITE, False, 0),
        (days,  11, False, OFFWHITE, False, 3),
    ], Inches(1.22), cy + Inches(0.12), Inches(1.75), ch - Inches(0.2))

    # Heading + body
    tb2([
        (head, 15, True, TEAL, False, 0),
        (body, 13, False, OFFWHITE, False, 5),
    ], Inches(3.2), cy + Inches(0.12), Inches(5.4), ch - Inches(0.15))

    cy += ch + gap

# ── "The Ask" banner ─────────────────────────────────────────────────────────
ask_y = cy + Inches(0.12)
rbox(Inches(1.18), ask_y, Inches(7.55), Inches(0.92), TEAL)

# Megaphone circle
oval(Inches(1.28), ask_y + Inches(0.12), Inches(0.68), Inches(0.68),
     RGBColor(0x0D, 0x2C, 0x2C))
tb("📣", Inches(1.28), ask_y + Inches(0.1), Inches(0.68), Inches(0.68),
   size=22, align=PP_ALIGN.CENTER)

tb("The Ask:", Inches(2.1), ask_y + Inches(0.15), Inches(1.3), Inches(0.45),
   size=16, bold=True, color=BG)

tb("Approve MVP build + small launch budget for content, ads, and the first mail drop.",
   Inches(3.45), ask_y + Inches(0.12), Inches(5.1), Inches(0.7),
   size=14, bold=True, color=BG)

# ── Right card panel ─────────────────────────────────────────────────────────
rx = Inches(9.18)
rw = Inches(3.75)
rh = Inches(6.85)
ry = Inches(0.35)

rbox(rx, ry, rw, rh, TEAL_CARD)

# Large teal circle
cr = Inches(1.6)
cx_c = rx + rw / 2
cy_c = ry + Inches(2.15)
oval(cx_c - cr, cy_c - cr, cr*2, cr*2, TEAL_DIM)

# Inner darker circle
cr2 = Inches(1.2)
oval(cx_c - cr2, cy_c - cr2, cr2*2, cr2*2, TEAL_CARD)

# Rocket emoji
tb("🚀", cx_c - Inches(0.9), cy_c - Inches(0.95), Inches(1.8), Inches(1.8),
   size=72, align=PP_ALIGN.CENTER)

# Divider line
box(rx + Inches(0.3), ry + Inches(4.55), rw - Inches(0.6), Inches(0.03),
    RGBColor(0x28, 0x60, 0x60))

# Contact info
tb("✉  dc.long123@gmail.com",
   rx + Inches(0.3), ry + Inches(4.75), rw - Inches(0.4), Inches(0.42),
   size=13, color=OFFWHITE)

tb("🌐  kidscreativityconfidenceclub.com",
   rx + Inches(0.3), ry + Inches(5.28), rw - Inches(0.4), Inches(0.42),
   size=13, color=OFFWHITE)

# Save
out = "/home/user/sales/K3C_Launch_Slide.pptx"
prs.save(out)
print(f"Saved: {out}")
